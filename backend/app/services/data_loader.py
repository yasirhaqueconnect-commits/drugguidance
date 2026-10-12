import os
import glob
import json
import openpyxl
from typing import List, Dict, Any

_backend_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
CHUNKS_CACHE_FILE = os.path.join(_backend_dir, "chunks_store.json")

def _find_files(pattern: str, base_dir: str) -> List[str]:
    """Finds matching dataset files in base_dir or datasets/ directory."""
    files = glob.glob(os.path.join(base_dir, pattern))
    if not files:
        files = glob.glob(os.path.join(base_dir, "datasets", pattern))
    if not files:
        root_data = os.path.abspath(os.path.join(_backend_dir, "..", "datasets"))
        files = glob.glob(os.path.join(root_data, pattern))
    return files

def load_all_excel_chunks(base_dir: str = None) -> List[Dict[str, Any]]:
    """
    Loads all 4 medical Excel files, extracts clean, rich text chunks,
    and returns a list of indexed chunk dictionaries.
    """
    if base_dir is None:
        base_dir = os.path.abspath(os.path.join(_backend_dir, ".."))
    
    chunks = []
    chunk_id = 1

    # Pre-scan / build Rx Lookup Dictionary from ALTERNATIVE MEDICINE
    rx_lookup: Dict[str, str] = {
        "panadol": "Non-Prescription (OTC)",
        "paracetamol": "Non-Prescription (OTC)",
        "calpol": "Non-Prescription (OTC)",
        "brufen": "Non-Prescription (OTC)",
        "ibuprofen": "Non-Prescription (OTC)",
        "disprin": "Non-Prescription (OTC)",
        "aspirin": "Non-Prescription (OTC)",
        "risek": "Non-Prescription (OTC)",
        "omeprazole": "Non-Prescription (OTC)",
        "zyrtec": "Non-Prescription (OTC)",
        "cetirizine": "Non-Prescription (OTC)",
        "benadryl": "Non-Prescription (OTC)",
        "piriton": "Non-Prescription (OTC)",
        "gravinate": "Non-Prescription (OTC)"
    }
    alt_files_scan = _find_files("*ALTERNATIVE*.xlsx", base_dir)
    for file_path in alt_files_scan:
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True)
            sheet = wb.active
            rows = list(sheet.iter_rows(values_only=True))
            if rows:
                for row in rows[1:]:
                    if not row or not any(row):
                        continue
                    m_name = str(row[0] or "").strip().lower()
                    g_name = str(row[1] or "").strip().lower()
                    rx_val = str(row[6] or "").strip().lower() if len(row) > 6 else ""
                    st = "Non-Prescription (OTC)" if rx_val in ["no", "otc", "otc/limited"] else "Prescription Required (Rx)"
                    if m_name:
                        rx_lookup[m_name] = st
                    if g_name:
                        rx_lookup[g_name] = st
        except Exception:
            pass

    # 1. Allergy-Based Drug Filtering
    allergy_files = _find_files("*ALLERGY*.xlsx", base_dir)
    for file_path in allergy_files:
        wb = openpyxl.load_workbook(file_path, read_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        # header = rows[0]
        for row in rows[1:]:
            if not row or not any(row):
                continue
            med_name = str(row[1] or "").strip()
            generic = str(row[2] or "").strip()
            chem_comp = str(row[3] or "").strip()
            active_mol = str(row[4] or "").strip()
            category = str(row[5] or "").strip()
            mfg = str(row[6] or "").strip()
            allergens = str(row[7] or "").strip()
            alt_med = str(row[8] or "").strip()
            remarks = str(row[9] or "").strip()

            if not med_name and not generic:
                continue

            rx_status = rx_lookup.get(med_name.lower(), rx_lookup.get(generic.lower(), "Prescription Required (Rx)"))

            text_content = (
                f"CATEGORY: Allergy & Sensitivity Guidance\n"
                f"Medicine Name: {med_name}\n"
                f"Generic / Active Ingredient: {generic}\n"
                f"Prescription Status: {rx_status}\n"
                f"Therapeutic Category: {category}\n"
                f"Manufacturer: {mfg}\n"
                f"Possible Allergic / Sensitizing Agents: {allergens}\n"
                f"Alternative Medicine (Allergen-Free): {alt_med}\n"
                f"Remarks: {remarks}"
            )
            kw = f"{med_name} {generic} {allergens} {alt_med} {category} {rx_status}".lower()
            if any(p in med_name.lower() or p in generic.lower() for p in ["panadol", "paracetamol", "calpol", "brufen", "ibuprofen"]):
                kw += " headache sar dard sir dard hedick fever bukhar pain bodyache"
            if any(bp in med_name.lower() or bp in generic.lower() for bp in ["amlodipine", "norvasc", "lisinopril", "zestril", "captopril", "capoten", "atenolol", "tenormin", "losartan", "cozaar"]):
                kw += " blood pressure bp hypertension high bp"

            chunks.append({
                "id": f"chunk_allergy_{chunk_id}",
                "category": "allergy",
                "medicine_name": med_name,
                "generic_name": generic,
                "keywords": kw,
                "content": text_content,
                "metadata": {
                    "source": "ALLERGY-BASED DRUG FILTERING",
                    "medicine": med_name,
                    "generic": generic,
                    "prescription_status": rx_status,
                    "allergens": allergens,
                    "alternatives": alt_med
                }
            })
            chunk_id += 1

    # 2. Alternative Medicine
    alt_files = _find_files("*ALTERNATIVE*.xlsx", base_dir)
    for file_path in alt_files:
        wb = openpyxl.load_workbook(file_path, read_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        for row in rows[1:]:
            if not row or not any(row):
                continue
            med_name = str(row[0] or "").strip()
            generic = str(row[1] or "").strip()
            form = str(row[2] or "").strip()
            mfg = str(row[3] or "").strip()
            price = str(row[4] or "").strip()
            alt_brands = str(row[5] or "").strip()
            rx_req = str(row[6] or "").strip()

            if not med_name and not generic:
                continue

            rx_status = "Non-Prescription (OTC)" if str(rx_req).strip().lower() in ["no", "otc", "otc/limited"] else "Prescription Required (Rx)"

            text_content = (
                f"CATEGORY: Alternative Medicine & Brand Substitutes\n"
                f"Medicine Name: {med_name}\n"
                f"Generic / Active Ingredient: {generic}\n"
                f"Prescription Status: {rx_status}\n"
                f"Dosage Form: {form}\n"
                f"Manufacturer: {mfg}\n"
                f"Price in PKR (approx): {price}\n"
                f"Alternative Brands / Substitutes: {alt_brands}\n"
                f"Prescription Required: {rx_req}"
            )
            kw = f"{med_name} {generic} {alt_brands} {form} {rx_status}".lower()
            if any(p in med_name.lower() or p in generic.lower() for p in ["panadol", "paracetamol", "calpol", "brufen", "ibuprofen"]):
                kw += " headache sar dard sir dard hedick fever bukhar pain bodyache"
            if any(bp in med_name.lower() or bp in generic.lower() for bp in ["amlodipine", "norvasc", "lisinopril", "zestril", "captopril", "capoten", "atenolol", "tenormin", "losartan", "cozaar"]):
                kw += " blood pressure bp hypertension high bp"

            chunks.append({
                "id": f"chunk_alt_{chunk_id}",
                "category": "alternative",
                "medicine_name": med_name,
                "generic_name": generic,
                "keywords": kw,
                "content": text_content,
                "metadata": {
                    "source": "ALTERNATIVE MEDICINE",
                    "medicine": med_name,
                    "generic": generic,
                    "prescription_status": rx_status,
                    "dosage_form": form,
                    "price_pkr": price,
                    "alternatives": alt_brands,
                    "prescription_required": rx_req
                }
            })
            chunk_id += 1

    # 3. Dosage Guidance
    dosage_files = _find_files("*Dosage*.xlsx", base_dir)
    for file_path in dosage_files:
        wb = openpyxl.load_workbook(file_path, read_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        for row in rows[1:]:
            if not row or not any(row):
                continue
            med_name = str(row[0] or "").strip()
            generic = str(row[1] or "").strip()
            form = str(row[2] or "").strip()
            strength = str(row[3] or "").strip()
            adult_dose = str(row[4] or "").strip()
            child_dose = str(row[5] or "").strip()
            max_dose = str(row[6] or "").strip()
            freq = str(row[7] or "").strip()

            if not med_name and not generic:
                continue

            rx_status = rx_lookup.get(med_name.lower(), rx_lookup.get(generic.lower(), "Prescription Required (Rx)"))

            dose_info_lines = []
            if adult_dose:
                dose_info_lines.append(f"Recommended Adult Dosage: {adult_dose}")
            if child_dose:
                dose_info_lines.append(f"Recommended Pediatric Dosage: {child_dose}")
            if max_dose:
                dose_info_lines.append(f"Maximum Daily Dose: {max_dose}")
            if freq:
                dose_info_lines.append(f"Timing / Frequency: {freq}")
            dose_info_str = "\n".join(dose_info_lines) if dose_info_lines else "Dosage: Follow standard therapeutic dosing under clinical guidance."

            text_content = (
                f"CATEGORY: Dosage & Formulation Reference\n"
                f"Medicine Name: {med_name}\n"
                f"Active Ingredient / Generic: {generic}\n"
                f"Prescription Status: {rx_status}\n"
                f"Available Dosage Form(s): {form}\n"
                f"Available Strength(s): {strength}\n"
                f"{dose_info_str}\n"
                f"Commercial Formulations: Available in {form} preparations with strengths of {strength}."
            )
            kw = f"{med_name} {generic} {form} {strength} {rx_status} dosage form strength formulation".lower()
            if any(p in med_name.lower() or p in generic.lower() for p in ["panadol", "paracetamol", "calpol", "brufen", "ibuprofen"]):
                kw += " headache sar dard sir dard hedick fever bukhar pain bodyache"
            if any(bp in med_name.lower() or bp in generic.lower() for bp in ["amlodipine", "norvasc", "lisinopril", "zestril", "captopril", "capoten", "atenolol", "tenormin", "losartan", "cozaar"]):
                kw += " blood pressure bp hypertension high bp"

            chunks.append({
                "id": f"chunk_dosage_{chunk_id}",
                "category": "dosage",
                "medicine_name": med_name,
                "generic_name": generic,
                "keywords": kw,
                "content": text_content,
                "metadata": {
                    "source": "Dosage Guidance",
                    "medicine": med_name,
                    "generic": generic,
                    "prescription_status": rx_status,
                    "dosage_form": form,
                    "strength": strength,
                    "adult_dosage": adult_dose,
                    "child_dosage": child_dose,
                    "max_daily_dose": max_dose,
                    "frequency": freq
                }
            })
            chunk_id += 1

    # 4. Drug-Drug Interactions
    interaction_files = _find_files("*INTERACTIONS*.xlsx", base_dir)
    for file_path in interaction_files:
        wb = openpyxl.load_workbook(file_path, read_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        for row in rows[1:]:
            if not row or not any(row):
                continue
            med1 = str(row[0] or "").strip()
            med2 = str(row[1] or "").strip()
            interaction_type = str(row[2] or "").strip()
            desc = str(row[3] or "").strip()
            clinical_effect = str(row[4] or "").strip()
            severity = str(row[5] or "").strip()
            guideline = str(row[6] or "").strip()

            if not med1 and not med2:
                continue

            text_content = (
                f"CATEGORY: Drug-Drug Interaction Warning\n"
                f"Medicine 1: {med1}\n"
                f"Medicine 2: {med2}\n"
                f"Severity Level: {severity}\n"
                f"Interaction Type: {interaction_type}\n"
                f"Interaction Description: {desc}\n"
                f"Clinical Effect / What Happens: {clinical_effect}\n"
                f"Clinical Recommendation / Guideline: {guideline}"
            )
            chunks.append({
                "id": f"chunk_interaction_{chunk_id}",
                "category": "interaction",
                "medicine_name": f"{med1} + {med2}",
                "generic_name": f"{med1}, {med2}",
                "keywords": f"{med1} {med2} {severity} {interaction_type} interaction".lower(),
                "content": text_content,
                "metadata": {
                    "source": "DRUG-DRUG INTERACTIONS",
                    "medicine_1": med1,
                    "medicine_2": med2,
                    "severity": severity,
                    "clinical_effect": clinical_effect,
                    "recommendation": guideline
                }
            })
            chunk_id += 1

    # Save to disk as fast local cache JSON
    with open(CHUNKS_CACHE_FILE, "w", encoding="utf-8") as f:
        json.dump(chunks, f, ensure_ascii=False, indent=2)

    print(f"Successfully processed and stored {len(chunks)} medical chunks into {CHUNKS_CACHE_FILE}")
    return chunks

def get_stored_chunks() -> List[Dict[str, Any]]:
    if os.path.exists(CHUNKS_CACHE_FILE):
        try:
            with open(CHUNKS_CACHE_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return load_all_excel_chunks()

if __name__ == "__main__":
    loaded = load_all_excel_chunks()
    print("Total chunks created:", len(loaded))
